import os

from django.shortcuts import get_object_or_404
from rest_framework.parsers import FormParser, MultiPartParser

from django.utils import timezone
from rest_framework import generics, status
from rest_framework.exceptions import PermissionDenied
from rest_framework.permissions import IsAuthenticated

from accounts.models import User
from accounts.permissions import IsAdmin, IsAdminOrManager
from etms.responses import success_response

from .filters import TaskFilter
from .models import Task, TaskActivity, TaskAttachment, TaskComment
from .serializers import (
    TaskActivitySerializer,
    TaskAttachmentSerializer,
    TaskCommentSerializer,
    TaskSerializer,
    TaskStatusSerializer,
    TaskWriteSerializer,
)


def log_task_activity(task, user, action, message):
    """Create an audit log record for a task action."""
    try:
        TaskActivity.objects.create(
            task=task,
            user=user if getattr(user, "is_authenticated", False) else None,
            action=action,
            message=message,
        )
    except Exception:
        pass


def visible_tasks_for(user):
    """The single place that decides which tasks a user is allowed to see."""
    tasks = Task.objects.select_related("assigned_to", "assigned_by")

    if user.role == User.Role.ADMIN:
        return tasks

    from django.db.models import Q
    profile = getattr(user, "employee_profile", None)  # None if no Employee record

    if user.role == User.Role.MANAGER:
        if profile is None:
            return tasks.filter(assigned_by=user)
        return tasks.filter(Q(assigned_to__manager=profile) | Q(assigned_by=user))

    # EMPLOYEE: tasks assigned to them OR tasks created by them
    if profile:
        return tasks.filter(Q(assigned_to=profile) | Q(assigned_by=user))
    return tasks.filter(assigned_by=user)


class BaseTaskListView(generics.ListAPIView):
    """Shared behaviour for all the task list endpoints."""

    serializer_class = TaskSerializer
    filterset_class = TaskFilter
    search_fields = ["title"]
    ordering_fields = ["deadline", "priority", "status", "created_at"]
    list_message = "Tasks fetched successfully"

    def get_queryset(self):
        return visible_tasks_for(self.request.user)

    def list(self, request, *args, **kwargs):
        response = super().list(request, *args, **kwargs)  # paginated + filtered
        return success_response(self.list_message, response.data)


class TaskListCreateView(BaseTaskListView):
    """GET /api/tasks/  and  POST /api/tasks/"""

    def get_permissions(self):
        if self.request.method == "POST":
            return [IsAuthenticated(), IsAdminOrManager()]
        return [IsAuthenticated()]

    def post(self, request, *args, **kwargs):
        if request.user.role == User.Role.EMPLOYEE:
            raise PermissionDenied("Employees are not permitted to create tasks.")

        serializer = TaskWriteSerializer(data=request.data, context={"request": request})
        serializer.is_valid(raise_exception=True)
        task = serializer.save(assigned_by=request.user)  # who created it is set to request.user

        # Audit log: task created
        actor_name = request.user.get_full_name() or request.user.username
        log_task_activity(task, request.user, "CREATED", f"Task created by {actor_name}")
        if task.assigned_to:
            assignee_name = task.assigned_to.full_name or "employee"
            log_task_activity(task, request.user, "ASSIGNED", f"Assigned to {assignee_name}")

        # Auto-create notification for assigned employee (only if assigned to someone else)
        try:
            if task.assigned_to and task.assigned_to.user and task.assigned_to.user != request.user:
                from notifications.models import Notification
                Notification.objects.create(
                    user=task.assigned_to.user,
                    title="New Task Assigned",
                    message=f"You have been assigned a new task: {task.title}",
                    notification_type=Notification.NotificationType.TASK_ASSIGNED,
                )
        except Exception:
            pass  # Never let notification failure break task creation

        return success_response("Task created successfully",
                                TaskSerializer(task).data, status.HTTP_201_CREATED)


class MyTasksView(BaseTaskListView):
    """GET /api/tasks/my-tasks/ - tasks assigned to me (any role that has an Employee profile)."""
    list_message = "My tasks fetched successfully"

    def get_queryset(self):
        profile = getattr(self.request.user, "employee_profile", None)
        tasks = Task.objects.select_related("assigned_to", "assigned_by")
        return tasks.filter(assigned_to=profile) if profile else tasks.none()


class OverdueTasksView(BaseTaskListView):
    """GET /api/tasks/overdue/"""
    list_message = "Overdue tasks fetched successfully"

    def get_queryset(self):
        return super().get_queryset().overdue()


class CompletedTasksView(BaseTaskListView):
    """GET /api/tasks/completed/"""
    list_message = "Completed tasks fetched successfully"

    def get_queryset(self):
        return super().get_queryset().completed()


class PendingTasksView(BaseTaskListView):
    """GET /api/tasks/pending/"""
    list_message = "Pending tasks fetched successfully"

    def get_queryset(self):
        return super().get_queryset().pending()


class TaskDetailView(generics.RetrieveUpdateDestroyAPIView):
    """GET / PUT / PATCH / DELETE  /api/tasks/<id>/"""

    serializer_class = TaskSerializer

    def get_queryset(self):
        # Unauthorised tasks simply don't exist for this user (404), which
        # also hides whether the task exists at all.
        return visible_tasks_for(self.request.user)

    def get_permissions(self):
        return [IsAuthenticated()]

    def retrieve(self, request, *args, **kwargs):
        return success_response("Task fetched successfully", TaskSerializer(self.get_object()).data)

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop("partial", False)
        task = self.get_object()

        if request.user.role == User.Role.EMPLOYEE:
            if set(request.data.keys()) - {"status"}:
                raise PermissionDenied("Employees can only update the task status.")
            serializer = TaskStatusSerializer(task, data=request.data, context={"request": request})
        else:
            serializer = TaskWriteSerializer(
                task, data=request.data, partial=partial, context={"request": request}
            )

        serializer.is_valid(raise_exception=True)

        # Capture old values before save
        old_status = task.status
        old_priority = task.priority
        old_assigned_to = task.assigned_to
        old_deadline = task.deadline
        old_title = task.title

        task = serializer.save()

        # Log activity trail
        actor_name = request.user.get_full_name() or request.user.username
        if old_status != task.status:
            log_task_activity(
                task, request.user, "STATUS_CHANGED",
                f"{actor_name} moved task to {task.get_status_display()}"
            )
        if old_assigned_to != task.assigned_to:
            new_name = task.assigned_to.full_name if task.assigned_to else "Unassigned"
            log_task_activity(
                task, request.user, "REASSIGNED",
                f"{actor_name} reassigned task to {new_name}"
            )
            # Reassignment resets status to PENDING so new employee accepts it
            if task.status != Task.Status.PENDING:
                task.status = Task.Status.PENDING
                task.accepted_at = None
                task.completed_at = None
                task.save(update_fields=["status", "accepted_at", "completed_at", "updated_at"])
        if old_deadline != task.deadline:
            log_task_activity(
                task, request.user, "DEADLINE_CHANGED",
                f"{actor_name} changed deadline to {task.deadline or 'None'}"
            )
        if old_priority != task.priority:
            log_task_activity(
                task, request.user, "PRIORITY_CHANGED",
                f"{actor_name} changed priority to {task.get_priority_display()}"
            )
        if old_title != task.title:
            log_task_activity(
                task, request.user, "TITLE_CHANGED",
                f"{actor_name} updated task title to \"{task.title}\""
            )


        # Auto-create notifications after update
        try:
            from notifications.models import Notification
            new_status = task.status

            # Notify creator when employee marks task completed
            if old_status != "COMPLETED" and new_status == "COMPLETED":
                if task.assigned_by and task.assigned_by != request.user:
                    Notification.objects.create(
                        user=task.assigned_by,
                        title="Task Completed ✅",
                        message=f"\"{task.title}\" has been marked as completed by "
                                f"{request.user.get_full_name() or request.user.email}.",
                        notification_type=Notification.NotificationType.TASK_COMPLETED,
                    )

            # Notify assigned employee when their task is updated by someone else
            elif task.assigned_to and task.assigned_to.user and task.assigned_to.user != request.user:
                Notification.objects.create(
                    user=task.assigned_to.user,
                    title="Task Updated",
                    message=f"Your task \"{task.title}\" has been updated by "
                            f"{request.user.get_full_name() or request.user.email}.",
                    notification_type=Notification.NotificationType.TASK_UPDATED,
                )
        except Exception:
            pass  # Never let notification failure break task update

        return success_response("Task updated successfully", TaskSerializer(task).data)

    def destroy(self, request, *args, **kwargs):
        task = self.get_object()
        is_creator = task.assigned_by_id == request.user.id
        if request.user.role == User.Role.EMPLOYEE and not is_creator and request.user.role != User.Role.ADMIN:
            raise PermissionDenied("You can only delete tasks you created.")
        task.delete()
        return success_response("Task deleted successfully", status_code=status.HTTP_204_NO_CONTENT)
    
    # ---------------------------------------------------------------
# Comments
# ---------------------------------------------------------------
class TaskCommentListCreateView(generics.ListCreateAPIView):
    """GET / POST  /api/tasks/<task_id>/comments/"""

    serializer_class = TaskCommentSerializer

    def get_task(self):
        # 404 if the task doesn't exist OR this user isn't allowed to see it.
        return get_object_or_404(visible_tasks_for(self.request.user), pk=self.kwargs["task_id"])

    def get_queryset(self):
        return TaskComment.objects.filter(task=self.get_task()).select_related("user")

    def list(self, request, *args, **kwargs):
        response = super().list(request, *args, **kwargs)
        return success_response("Comments fetched successfully", response.data)

    def create(self, request, *args, **kwargs):
        task = self.get_task()
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        comment = serializer.save(task=task, user=request.user)

        # Audit log for comment
        actor_name = request.user.get_full_name() or request.user.username
        comment_snippet = comment.comment[:50] + ("..." if len(comment.comment) > 50 else "")
        log_task_activity(task, request.user, "COMMENT", f"{actor_name} commented: \"{comment_snippet}\"")

        # Notify the other party about the new comment
        try:
            from notifications.models import Notification
            recipient = None
            if task.assigned_to and task.assigned_to.user == request.user:
                if task.assigned_by and task.assigned_by != request.user:
                    recipient = task.assigned_by
            else:
                if task.assigned_to and task.assigned_to.user and task.assigned_to.user != request.user:
                    recipient = task.assigned_to.user

            if recipient:
                sender_name = request.user.get_full_name() or request.user.username
                comment_preview = comment.comment[:80] + ("..." if len(comment.comment) > 80 else "")
                Notification.objects.create(
                    user=recipient,
                    title=f"New comment on: {task.title}",
                    message=f"{sender_name}: \"{comment_preview}\"",
                    notification_type=Notification.NotificationType.TASK_UPDATED,
                )
        except Exception:
            pass  # Never let notification failure break comment creation

        return success_response("Comment added successfully",
                                TaskCommentSerializer(comment).data, status.HTTP_201_CREATED)


class CommentDetailView(generics.RetrieveUpdateDestroyAPIView):
    """PUT / DELETE  /api/comments/<id>/  (PATCH also works; GET is switched off)"""

    serializer_class = TaskCommentSerializer
    http_method_names = ["put", "patch", "delete", "options"]

    def get_queryset(self):
        # Only comments on tasks this user is allowed to see.
        return TaskComment.objects.filter(
            task__in=visible_tasks_for(self.request.user)
        ).select_related("user")

    def update(self, request, *args, **kwargs):
        comment = self.get_object()
        if comment.user_id != request.user.id:
            raise PermissionDenied("You can only edit your own comments.")
        partial = kwargs.pop("partial", False)
        serializer = self.get_serializer(comment, data=request.data, partial=partial)
        serializer.is_valid(raise_exception=True)
        comment = serializer.save()
        return success_response("Comment updated successfully", TaskCommentSerializer(comment).data)

    def destroy(self, request, *args, **kwargs):
        comment = self.get_object()
        if comment.user_id != request.user.id and request.user.role != User.Role.ADMIN:
            raise PermissionDenied("You can only delete your own comments.")
        comment.delete()
        return success_response("Comment deleted successfully", status_code=status.HTTP_204_NO_CONTENT)


# ---------------------------------------------------------------
# Attachments
# ---------------------------------------------------------------
class TaskAttachmentListCreateView(generics.ListCreateAPIView):
    """GET / POST  /api/tasks/<task_id>/attachments/"""

    serializer_class = TaskAttachmentSerializer
    parser_classes = [MultiPartParser, FormParser]  # needed for file uploads

    def get_task(self):
        return get_object_or_404(visible_tasks_for(self.request.user), pk=self.kwargs["task_id"])

    def get_queryset(self):
        return TaskAttachment.objects.filter(task=self.get_task()).select_related("uploaded_by")

    def list(self, request, *args, **kwargs):
        response = super().list(request, *args, **kwargs)
        return success_response("Attachments fetched successfully", response.data)

    def create(self, request, *args, **kwargs):
        task = self.get_task()
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)  # runs the type + size validator
        original_name = os.path.basename(serializer.validated_data["file"].name)
        attachment = serializer.save(task=task, uploaded_by=request.user, file_name=original_name[:255])

        # Audit log for attachment
        actor_name = request.user.get_full_name() or request.user.username
        log_task_activity(task, request.user, "ATTACHMENT", f"{actor_name} attached file: {original_name}")

        return success_response(
            "Attachment uploaded successfully",
            self.get_serializer(attachment).data,
            status.HTTP_201_CREATED,
        )


class AttachmentDeleteView(generics.DestroyAPIView):
    """DELETE /api/attachments/<id>/"""

    def get_queryset(self):
        return TaskAttachment.objects.filter(task__in=visible_tasks_for(self.request.user))

    def destroy(self, request, *args, **kwargs):
        attachment = self.get_object()
        if attachment.uploaded_by_id != request.user.id and request.user.role != User.Role.ADMIN:
            raise PermissionDenied("You can only delete attachments you uploaded.")
        attachment.file.delete(save=False)  # remove the actual file from the media folder
        attachment.delete()                 # then remove the database row
        return success_response("Attachment deleted successfully", status_code=status.HTTP_204_NO_CONTENT)


# ---------------------------------------------------------------
# Activities / Audit Trail
# ---------------------------------------------------------------
class TaskActivityListView(generics.ListAPIView):
    """GET /api/tasks/<task_id>/activities/"""

    serializer_class = TaskActivitySerializer
    permission_classes = [IsAuthenticated]

    def get_task(self):
        return get_object_or_404(visible_tasks_for(self.request.user), pk=self.kwargs["task_id"])

    def get_queryset(self):
        return TaskActivity.objects.filter(task=self.get_task()).select_related("user")

    def list(self, request, *args, **kwargs):
        response = super().list(request, *args, **kwargs)
        return success_response("Task activities fetched successfully", response.data)


class TaskAcceptView(generics.GenericAPIView):
    """
    POST / PATCH /api/tasks/<pk>/accept/
    Allows the assigned employee to accept the task.
    Automatically moves task status from PENDING to IN_PROGRESS.
    """
    permission_classes = [IsAuthenticated]

    def get_task(self):
        return get_object_or_404(Task.objects.select_related("assigned_to__user", "assigned_by"), pk=self.kwargs["pk"])

    def patch(self, request, *args, **kwargs):
        return self._process_accept(request, *args, **kwargs)

    def post(self, request, *args, **kwargs):
        return self._process_accept(request, *args, **kwargs)

    def _process_accept(self, request, *args, **kwargs):
        from rest_framework.response import Response
        task = self.get_task()
        user = request.user

        # Validation: ONLY the assigned employee can accept the task
        if user.role != User.Role.EMPLOYEE:
            raise PermissionDenied("Only the assigned employee can accept this task.")

        is_assigned_user = (task.assigned_to and task.assigned_to.user_id == user.id)
        if not is_assigned_user:
            raise PermissionDenied("Only the assigned employee can accept this task.")

        # 400 if the task status is not "pending"
        if task.status != Task.Status.PENDING:
            return Response(
                {"success": False, "message": f"Task cannot be accepted because its status is already {task.get_status_display()}."},
                status=status.HTTP_400_BAD_REQUEST
            )

        task.status = Task.Status.IN_PROGRESS
        task.accepted_at = timezone.now()
        task.save(update_fields=["status", "accepted_at", "updated_at"])

        actor_name = user.get_full_name() or user.username
        log_task_activity(
            task, user, "STATUS_CHANGED",
            f"{actor_name} accepted the task (Status: In Progress)"
        )

        # Notify task creator / manager: "<Employee name> accepted the task: <title>"
        try:
            from notifications.models import Notification
            if task.assigned_by:
                Notification.objects.create(
                    user=task.assigned_by,
                    title="Task Accepted 🚀",
                    message=f"{actor_name} accepted the task: {task.title}",
                    notification_type=Notification.NotificationType.TASK_UPDATED,
                )
        except Exception:
            pass

        return success_response("Task accepted and moved to In Progress", TaskSerializer(task).data)


class TaskRejectView(generics.GenericAPIView):
    """
    POST / PATCH /api/tasks/<pk>/reject/
    Body: { "reason": "<mandatory reason, min 10 chars>" }
    Allows the assigned employee to reject a pending task.
    Automatically moves task status from PENDING to REJECTED.
    """
    permission_classes = [IsAuthenticated]

    def get_task(self):
        return get_object_or_404(Task.objects.select_related("assigned_to__user", "assigned_by"), pk=self.kwargs["pk"])

    def patch(self, request, *args, **kwargs):
        return self._process_reject(request, *args, **kwargs)

    def post(self, request, *args, **kwargs):
        return self._process_reject(request, *args, **kwargs)

    def _process_reject(self, request, *args, **kwargs):
        from rest_framework.response import Response
        task = self.get_task()
        user = request.user

        if user.role != User.Role.EMPLOYEE:
            raise PermissionDenied("Only the assigned employee can reject this task.")

        is_assigned_user = (task.assigned_to and task.assigned_to.user_id == user.id)
        if not is_assigned_user:
            raise PermissionDenied("Only the assigned employee can reject this task.")

        if task.status != Task.Status.PENDING:
            return Response(
                {"success": False, "message": f"Task cannot be rejected because its status is already {task.get_status_display()}."},
                status=status.HTTP_400_BAD_REQUEST
            )

        reason = request.data.get("reason", "")
        if not isinstance(reason, str):
            reason = ""
        reason = reason.strip()
        if not reason or len(reason) < 10:
            return Response(
                {"success": False, "message": "A reason of at least 10 characters is required to reject a task."},
                status=status.HTTP_400_BAD_REQUEST
            )

        task.status = Task.Status.REJECTED
        task.status_reason = reason
        task.rejected_at = timezone.now()
        task.save(update_fields=["status", "status_reason", "rejected_at", "updated_at"])

        actor_name = user.get_full_name() or user.username
        msg = f"{actor_name} rejected the task — Reason: {reason}"
        log_task_activity(task, user, "STATUS_CHANGED", msg)

        # Notify task creator / manager, including the reason
        try:
            from notifications.models import Notification
            if task.assigned_by:
                Notification.objects.create(
                    user=task.assigned_by,
                    title="Task Rejected ❌",
                    message=f'"{task.title}" was rejected by {actor_name}. Reason: {reason}',
                    notification_type=Notification.NotificationType.TASK_UPDATED,
                )
        except Exception:
            pass

        return success_response("Task rejected successfully", TaskSerializer(task).data)


class TaskAbortView(generics.GenericAPIView):
    """
    POST / PATCH /api/tasks/<pk>/abort/
    Body: { "reason": "<mandatory reason, min 10 chars>" }
    Allows the assigned employee to abort an in-progress task.
    Automatically moves task status from IN_PROGRESS to ABORTED.
    """
    permission_classes = [IsAuthenticated]

    def get_task(self):
        return get_object_or_404(Task.objects.select_related("assigned_to__user", "assigned_by"), pk=self.kwargs["pk"])

    def patch(self, request, *args, **kwargs):
        return self._process_abort(request, *args, **kwargs)

    def post(self, request, *args, **kwargs):
        return self._process_abort(request, *args, **kwargs)

    def _process_abort(self, request, *args, **kwargs):
        from rest_framework.response import Response
        task = self.get_task()
        user = request.user

        if user.role != User.Role.EMPLOYEE:
            raise PermissionDenied("Only the assigned employee can abort this task.")

        is_assigned_user = (task.assigned_to and task.assigned_to.user_id == user.id)
        if not is_assigned_user:
            raise PermissionDenied("Only the assigned employee can abort this task.")

        if task.status != Task.Status.IN_PROGRESS:
            return Response(
                {"success": False, "message": f"Task cannot be aborted because its status is {task.get_status_display()} (only In Progress tasks can be aborted)."},
                status=status.HTTP_400_BAD_REQUEST
            )

        reason = request.data.get("reason", "")
        if not isinstance(reason, str):
            reason = ""
        reason = reason.strip()
        if not reason or len(reason) < 10:
            return Response(
                {"success": False, "message": "A reason of at least 10 characters is required to abort a task."},
                status=status.HTTP_400_BAD_REQUEST
            )

        task.status = Task.Status.ABORTED
        task.status_reason = reason
        task.aborted_at = timezone.now()
        task.save(update_fields=["status", "status_reason", "aborted_at", "updated_at"])

        actor_name = user.get_full_name() or user.username
        msg = f"{actor_name} aborted the task — Reason: {reason}"
        log_task_activity(task, user, "STATUS_CHANGED", msg)

        # Notify task creator / manager, including the reason
        try:
            from notifications.models import Notification
            if task.assigned_by:
                Notification.objects.create(
                    user=task.assigned_by,
                    title="Task Aborted ⚠️",
                    message=f'"{task.title}" was aborted by {actor_name}. Reason: {reason}',
                    notification_type=Notification.NotificationType.TASK_UPDATED,
                )
        except Exception:
            pass

        return success_response("Task aborted successfully", TaskSerializer(task).data)


class TaskCompleteView(generics.GenericAPIView):
    """
    POST / PATCH /api/tasks/<pk>/complete/
    Allows the assigned employee to mark the task completed.
    Automatically moves task status to COMPLETED and sets completed_at.
    """
    permission_classes = [IsAuthenticated]

    def get_task(self):
        return get_object_or_404(Task.objects.select_related("assigned_to__user", "assigned_by"), pk=self.kwargs["pk"])

    def patch(self, request, *args, **kwargs):
        return self._process_complete(request, *args, **kwargs)

    def post(self, request, *args, **kwargs):
        return self._process_complete(request, *args, **kwargs)

    def _process_complete(self, request, *args, **kwargs):
        from rest_framework.response import Response
        task = self.get_task()
        user = request.user

        if user.role != User.Role.EMPLOYEE:
            raise PermissionDenied("Only the assigned employee can mark this task completed.")

        is_assigned_user = (task.assigned_to and task.assigned_to.user_id == user.id)
        if not is_assigned_user:
            raise PermissionDenied("Only the assigned employee can mark this task completed.")

        # Locked tasks or invalid state
        if task.status in (Task.Status.COMPLETED, Task.Status.REJECTED, Task.Status.ABORTED, Task.Status.CANCELLED):
            return Response(
                {"success": False, "message": f"This task is locked ({task.get_status_display()}) and cannot be modified."},
                status=status.HTTP_400_BAD_REQUEST
            )

        if task.status != Task.Status.IN_PROGRESS:
            return Response(
                {"success": False, "message": f"Only In Progress tasks can be marked completed (current status: {task.get_status_display()})."},
                status=status.HTTP_400_BAD_REQUEST
            )

        task.status = Task.Status.COMPLETED
        task.completed_at = timezone.now()
        task.save(update_fields=["status", "completed_at", "updated_at"])

        actor_name = user.get_full_name() or user.username
        log_task_activity(
            task, user, "STATUS_CHANGED",
            f"{actor_name} completed the task ✅"
        )

        # Notify task creator / manager
        try:
            from notifications.models import Notification
            if task.assigned_by and task.assigned_by != user:
                Notification.objects.create(
                    user=task.assigned_by,
                    title="Task Completed ✅",
                    message=f'"{task.title}" has been marked as completed by {actor_name}.',
                    notification_type=Notification.NotificationType.TASK_COMPLETED,
                )
        except Exception:
            pass

        return success_response("Task marked as completed", TaskSerializer(task).data)


class TaskReassignView(generics.GenericAPIView):
    """
    PATCH /api/tasks/<pk>/reassign/
    Body: { "assigned_to": <employee_id (pk)>, "note": "<optional reason>" }
    Only ADMIN and MANAGER can reassign tasks.
    Manager can only reassign to their own team members.
    """
    permission_classes = [IsAdminOrManager]

    def get_task(self):
        return get_object_or_404(visible_tasks_for(self.request.user), pk=self.kwargs["pk"])

    def patch(self, request, *args, **kwargs):
        from employees.models import Employee

        task = self.get_task()
        user = request.user

        new_assignee_id = request.data.get("assigned_to")
        note = request.data.get("note", "").strip()

        if not new_assignee_id:
            from rest_framework.response import Response
            return Response(
                {"message": "assigned_to field is required."},
                status=status.HTTP_400_BAD_REQUEST
            )

        # Fetch new employee
        try:
            new_employee = Employee.objects.select_related("user").get(pk=new_assignee_id)
        except Employee.DoesNotExist:
            from rest_framework.response import Response
            return Response(
                {"message": "Employee not found."},
                status=status.HTTP_404_NOT_FOUND
            )

        if not new_employee.is_active:
            from rest_framework.response import Response
            return Response(
                {"message": "Cannot assign a task to an inactive employee."},
                status=status.HTTP_400_BAD_REQUEST
            )

        # Manager can only reassign to employees (cannot reassign to other managers or admins)
        if user.role == User.Role.MANAGER:
            manager_profile = getattr(user, "employee_profile", None)
            if manager_profile is None:
                from employees.models import Employee
                from django.utils import timezone
                manager_profile, _ = Employee.objects.get_or_create(
                    user=user,
                    defaults={
                        "employee_id": f"EMP{user.id:04d}",
                        "full_name": user.get_full_name() or user.username,
                        "email": user.email or f"{user.username}@company.com",
                        "department": "Management",
                        "designation": "Manager",
                        "joining_date": timezone.localdate(),
                    }
                )
            # Manager cannot assign to other managers or admins
            if new_employee.user and new_employee.user.role in (User.Role.ADMIN, User.Role.MANAGER) and new_employee.pk != manager_profile.pk:
                from rest_framework.response import Response
                return Response(
                    {"message": "Managers can only assign tasks to employees."},
                    status=status.HTTP_400_BAD_REQUEST
                )

            # Manager can only reassign to employees within their team
            if new_employee.manager_id != manager_profile.pk and new_employee.pk != manager_profile.pk:
                from rest_framework.response import Response
                return Response(
                    {"message": "You can only assign tasks to employees in your team."},
                    status=status.HTTP_400_BAD_REQUEST
                )

        old_employee = task.assigned_to
        old_name = old_employee.full_name if old_employee else "Unassigned"
        new_name = new_employee.full_name
        actor_name = user.get_full_name() or user.username

        # Save (Reset status to PENDING so new assignee must accept it)
        task.assigned_to = new_employee
        task.status = Task.Status.PENDING
        task.accepted_at = None
        task.completed_at = None
        task.save(update_fields=["assigned_to", "status", "accepted_at", "completed_at", "updated_at"])

        # Audit log
        msg = f"{actor_name} reassigned task from {old_name} to {new_name}"
        if note:
            msg += f" — Reason: {note}"
        log_task_activity(task, user, "REASSIGNED", msg)

        # Notify the new assignee
        try:
            from notifications.models import Notification
            if new_employee.user and new_employee.user != user:
                Notification.objects.create(
                    user=new_employee.user,
                    title="Task Reassigned to You 🔄",
                    message=(
                        f'"{task.title}" has been reassigned to you by {actor_name}.'
                        + (f" Note: {note}" if note else "")
                    ),
                    notification_type=Notification.NotificationType.TASK_ASSIGNED,
                )
            # Also notify old assignee (if different from actor and new assignee)
            if (old_employee and old_employee.user and
                    old_employee.user != user and old_employee != new_employee):
                Notification.objects.create(
                    user=old_employee.user,
                    title="Task Reassigned",
                    message=(
                        f'"{task.title}" has been reassigned from you to {new_name} by {actor_name}.'
                    ),
                    notification_type=Notification.NotificationType.TASK_UPDATED,
                )
        except Exception:
            pass

        return success_response("Task reassigned successfully", TaskSerializer(task).data)


# ---------------------------------------------------------------
# Exports — CSV / Excel / PDF
# ---------------------------------------------------------------
import csv
import io
from datetime import date

from django.http import HttpResponse
from rest_framework.views import APIView


# ── Shared helper ────────────────────────────────────────────────────────────
EXPORT_HEADERS = [
    "Task ID", "Title", "Description", "Category",
    "Priority", "Status",
    "Assigned To", "Assigned By",
    "Start Date", "Deadline", "Completed At", "Created At",
]


def _get_filtered_qs(request):
    """Return the role-filtered queryset with optional query-param filters applied."""
    qs = visible_tasks_for(request.user).select_related("assigned_to__user", "assigned_by")

    status_p   = request.query_params.get("status")
    priority_p = request.query_params.get("priority")
    search_p   = request.query_params.get("search")

    if status_p and status_p.upper() in [c[0] for c in Task.Status.choices]:
        qs = qs.filter(status=status_p.upper())
    if priority_p and priority_p.upper() in [c[0] for c in Task.Priority.choices]:
        qs = qs.filter(priority=priority_p.upper())
    if search_p:
        qs = qs.filter(title__icontains=search_p)

    return qs


def _task_row(task):
    """Return a list of values for one task — used by all three exporters."""
    assigned_to = ""
    if task.assigned_to:
        assigned_to = task.assigned_to.full_name or getattr(
            task.assigned_to.user, "get_full_name", lambda: ""
        )()

    assigned_by = ""
    if task.assigned_by:
        assigned_by = task.assigned_by.get_full_name() or task.assigned_by.username

    return [
        task.pk,
        task.title,
        task.description,
        task.category,
        task.get_priority_display(),
        task.get_status_display(),
        assigned_to,
        assigned_by,
        str(task.start_date) if task.start_date else "",
        str(task.deadline) if task.deadline else "",
        task.completed_at.strftime("%Y-%m-%d %H:%M") if task.completed_at else "",
        task.created_at.strftime("%Y-%m-%d %H:%M"),
    ]


def _today_str():
    return date.today().strftime("%Y-%m-%d")


# ── CSV ───────────────────────────────────────────────────────────────────────
class TaskExportCSVView(APIView):
    """GET /api/tasks/export/csv/"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        qs = _get_filtered_qs(request)
        response = HttpResponse(content_type="text/csv")
        response["Content-Disposition"] = f'attachment; filename="ETMS_Tasks_{_today_str()}.csv"'

        writer = csv.writer(response)
        writer.writerow(EXPORT_HEADERS)
        for task in qs:
            writer.writerow(_task_row(task))

        return response


# ── Excel (.xlsx) ─────────────────────────────────────────────────────────────
class TaskExportExcelView(APIView):
    """GET /api/tasks/export/excel/"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        except ImportError:
            return HttpResponse("openpyxl not installed", status=500)

        qs = _get_filtered_qs(request)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Tasks"

        # ── Header row styling ─────────────────────────────────────────────
        header_font  = Font(bold=True, color="FFFFFF", size=11)
        header_fill  = PatternFill("solid", fgColor="1E3A5F")
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left_align   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
        thin_border  = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"), bottom=Side(style="thin"),
        )

        # Priority colours
        priority_colors = {"High": "FFD6D6", "Medium": "FFF3CD", "Low": "D4EDDA"}
        status_colors   = {
            "Pending": "FFF3CD", "In Progress": "CCE5FF",
            "Completed": "D4EDDA", "Cancelled": "E2E3E5",
        }

        col_widths = [10, 30, 40, 15, 12, 14, 22, 22, 13, 13, 18, 18]
        for i, (header, width) in enumerate(zip(EXPORT_HEADERS, col_widths), start=1):
            cell = ws.cell(row=1, column=i, value=header)
            cell.font      = header_font
            cell.fill      = header_fill
            cell.alignment = center_align
            cell.border    = thin_border
            ws.column_dimensions[cell.column_letter].width = width

        ws.row_dimensions[1].height = 22

        # ── Data rows ──────────────────────────────────────────────────────
        for row_idx, task in enumerate(qs, start=2):
            row_data = _task_row(task)
            priority_label = task.get_priority_display()
            status_label   = task.get_status_display()

            for col_idx, value in enumerate(row_data, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border    = thin_border
                cell.alignment = left_align

                # Colour-code priority column (col 5) and status column (col 6)
                if col_idx == 5 and priority_label in priority_colors:
                    cell.fill = PatternFill("solid", fgColor=priority_colors[priority_label])
                if col_idx == 6 and status_label in status_colors:
                    cell.fill = PatternFill("solid", fgColor=status_colors[status_label])

            ws.row_dimensions[row_idx].height = 18

        # ── Freeze header row ──────────────────────────────────────────────
        ws.freeze_panes = "A2"

        # ── Auto-filter ────────────────────────────────────────────────────
        ws.auto_filter.ref = ws.dimensions

        # ── Stream to response ─────────────────────────────────────────────
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        response = HttpResponse(
            buf.read(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="ETMS_Tasks_{_today_str()}.xlsx"'
        return response


# ── PDF ───────────────────────────────────────────────────────────────────────
class TaskExportPDFView(APIView):
    """GET /api/tasks/export/pdf/"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        try:
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib import colors
            from reportlab.lib.units import cm
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.platypus import (
                SimpleDocTemplate, Table, TableStyle,
                Paragraph, Spacer, HRFlowable,
            )
            from reportlab.lib.enums import TA_CENTER, TA_LEFT
        except ImportError:
            return HttpResponse("reportlab not installed", status=500)

        qs   = _get_filtered_qs(request)
        user = request.user

        buf = io.BytesIO()
        doc = SimpleDocTemplate(
            buf,
            pagesize=landscape(A4),
            leftMargin=1.5*cm, rightMargin=1.5*cm,
            topMargin=1.5*cm, bottomMargin=1.5*cm,
            title="ETMS Task Report",
        )

        styles = getSampleStyleSheet()
        BLACK  = colors.black
        NAVY   = colors.HexColor("#1E3A5F")
        LGRAY  = colors.HexColor("#F2F2F2")
        MGRAY  = colors.HexColor("#CCCCCC")
        WHITE  = colors.white

        # Colour maps
        priority_map = {
            "High": colors.HexColor("#FFD6D6"),
            "Medium": colors.HexColor("#FFF3CD"),
            "Low": colors.HexColor("#D4EDDA"),
        }
        status_map = {
            "Pending":     colors.HexColor("#FFF3CD"),
            "In Progress": colors.HexColor("#CCE5FF"),
            "Completed":   colors.HexColor("#D4EDDA"),
            "Cancelled":   colors.HexColor("#E2E3E5"),
        }

        title_style = ParagraphStyle(
            "title", fontName="Helvetica-Bold", fontSize=16,
            textColor=NAVY, spaceAfter=4, alignment=TA_LEFT,
        )
        sub_style = ParagraphStyle(
            "sub", fontName="Helvetica", fontSize=9,
            textColor=colors.HexColor("#555555"), spaceAfter=2,
        )
        cell_style = ParagraphStyle(
            "cell", fontName="Helvetica", fontSize=8, leading=11,
        )
        hdr_style = ParagraphStyle(
            "hdr", fontName="Helvetica-Bold", fontSize=8,
            textColor=WHITE, leading=11, alignment=TA_CENTER,
        )

        story = []

        # ── Title block ────────────────────────────────────────────────────
        story.append(Paragraph("ETMS — Task Report", title_style))
        story.append(Paragraph(
            f"Generated: {_today_str()}  |  User: {user.get_full_name() or user.username}  |  Role: {user.role}",
            sub_style,
        ))
        story.append(Spacer(1, 0.15*cm))
        story.append(HRFlowable(width="100%", thickness=1.5, color=NAVY))
        story.append(Spacer(1, 0.3*cm))

        if not qs.exists():
            story.append(Paragraph("No tasks found for the selected filters.", styles["Normal"]))
        else:
            # ── Table headers ──────────────────────────────────────────────
            SHORT_HEADERS = [
                "#", "Title", "Category", "Priority", "Status",
                "Assigned To", "Assigned By", "Deadline", "Created",
            ]
            col_widths_pdf = [
                1.0*cm, 6.5*cm, 3.0*cm, 2.2*cm, 2.5*cm,
                4.0*cm, 4.0*cm, 2.8*cm, 2.8*cm,
            ]

            table_data = [[Paragraph(h, hdr_style) for h in SHORT_HEADERS]]

            for task in qs:
                row = _task_row(task)
                priority_label = task.get_priority_display()
                status_label   = task.get_status_display()

                table_data.append([
                    Paragraph(str(row[0]), cell_style),   # ID
                    Paragraph(row[1][:60], cell_style),   # Title
                    Paragraph(row[3], cell_style),        # Category
                    Paragraph(priority_label, cell_style),
                    Paragraph(status_label, cell_style),
                    Paragraph(row[6][:30], cell_style),   # Assigned To
                    Paragraph(row[7][:30], cell_style),   # Assigned By
                    Paragraph(row[9], cell_style),        # Deadline
                    Paragraph(row[11][:10], cell_style),  # Created At
                ])

            t = Table(table_data, colWidths=col_widths_pdf, repeatRows=1)

            # Base style
            ts = TableStyle([
                # Header
                ("BACKGROUND",    (0, 0), (-1, 0), NAVY),
                ("TEXTCOLOR",     (0, 0), (-1, 0), WHITE),
                ("FONTNAME",      (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE",      (0, 0), (-1, 0), 8),
                ("ALIGN",         (0, 0), (-1, 0), "CENTER"),
                # Body
                ("FONTSIZE",      (0, 1), (-1, -1), 8),
                ("VALIGN",        (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING",    (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ("LEFTPADDING",   (0, 0), (-1, -1), 4),
                ("GRID",          (0, 0), (-1, -1), 0.4, MGRAY),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [WHITE, LGRAY]),
            ])

            # Per-row colour for Priority & Status columns (col 3 & 4)
            for i, task in enumerate(qs, start=1):
                p_color = priority_map.get(task.get_priority_display())
                s_color = status_map.get(task.get_status_display())
                if p_color:
                    ts.add("BACKGROUND", (3, i), (3, i), p_color)
                if s_color:
                    ts.add("BACKGROUND", (4, i), (4, i), s_color)

            t.setStyle(ts)
            story.append(t)

        # ── Footer note ────────────────────────────────────────────────────
        story.append(Spacer(1, 0.4*cm))
        story.append(HRFlowable(width="100%", thickness=0.6, color=MGRAY))
        story.append(Spacer(1, 0.1*cm))
        story.append(Paragraph(
            f"Employee Task Management System  ·  {_today_str()}  ·  Confidential",
            ParagraphStyle("footer", fontName="Helvetica-Oblique", fontSize=7,
                           textColor=colors.HexColor("#888888"), alignment=TA_CENTER),
        ))

        doc.build(story)
        buf.seek(0)

        response = HttpResponse(buf.read(), content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="ETMS_Tasks_{_today_str()}.pdf"'
        return response